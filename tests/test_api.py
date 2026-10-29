"""End-to-end tests for the DocFlow API.

Uses a throwaway data dir + SQLite DB (env set before app import).
Run: .venv/bin/python -m pytest tests/ -q
"""
import io
import os
import tempfile
import time

_TMP = tempfile.mkdtemp(prefix="docflow-test-")
os.environ["DOCFLOW_DATA_DIR"] = _TMP
os.environ["DOCFLOW_DATABASE_URL"] = f"sqlite:///{_TMP}/test.db"

from fastapi.testclient import TestClient  # noqa: E402

from app.main import app  # noqa: E402


def wait_for(client: TestClient, job_id: str, want: set[str], timeout: float = 60.0) -> dict:
    deadline = time.time() + timeout
    last = {}
    while time.time() < deadline:
        resp = client.get(f"/jobs/{job_id}")
        assert resp.status_code == 200
        last = resp.json()
        if last["status"] in want:
            return last
        time.sleep(0.25)
    raise AssertionError(f"job {job_id} did not reach {want}; last={last}")


def test_health():
    with TestClient(app) as client:
        resp = client.get("/health")
        assert resp.status_code == 200
        assert resp.json()["status"] == "ok"


def test_full_chain_txt_upload_process_result():
    body = (
        "Quarterly Invoice Summary\n"
        "Please remit $1,250.00 by 2026-11-15.\n"
        "Questions: billing@example.com or ops@example.com\n"
        "This document summarizes the quarterly invoice for services rendered.\n"
    )
    with TestClient(app) as client:
        resp = client.post(
            "/jobs", files={"file": ("invoice.txt", io.BytesIO(body.encode()), "text/plain")}
        )
        assert resp.status_code == 202
        job_id = resp.json()["id"]

        job = wait_for(client, job_id, {"done", "failed"})
        assert job["status"] == "done", job
        analysis = job["result"]["analysis"]
        assert analysis["backend"] == "rules"
        assert analysis["title"] == "Quarterly Invoice Summary"
        assert "billing@example.com" in analysis["emails"]
        assert any("1,250.00" in a for a in analysis["amounts"])
        assert "2026-11-15" in analysis["dates"]
        assert analysis["word_count"] > 10


def test_docx_chain_converts_to_pdf():
    import docx

    doc = docx.Document()
    doc.add_heading("Service Agreement", level=1)
    doc.add_paragraph("Total fee is $3,400.00. Contact legal@example.com.")
    buf = io.BytesIO()
    doc.save(buf)

    with TestClient(app) as client:
        resp = client.post(
            "/jobs",
            files={
                "file": (
                    "agreement.docx",
                    io.BytesIO(buf.getvalue()),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        assert resp.status_code == 202
        job = wait_for(client, resp.json()["id"], {"done", "failed"}, timeout=120)
        assert job["status"] == "done", job
        result = job["result"]
        assert result["analysis"]["title"] == "Service Agreement"
        # LibreOffice is present on this host, so DOCX->PDF must have run.
        assert result["converted_file"] and result["converted_file"].endswith(".pdf")


def test_unsupported_file_rejected():
    with TestClient(app) as client:
        resp = client.post(
            "/jobs", files={"file": ("malware.exe", io.BytesIO(b"MZ..."), "application/octet-stream")}
        )
        assert resp.status_code == 400
        assert "unsupported" in resp.json()["detail"]


def test_job_not_found():
    with TestClient(app) as client:
        resp = client.get("/jobs/doesnotexist")
        assert resp.status_code == 404


def test_corrupt_pdf_fails_with_error():
    with TestClient(app) as client:
        resp = client.post(
            "/jobs",
            files={"file": ("broken.pdf", io.BytesIO(b"this is not a pdf at all"), "application/pdf")},
        )
        assert resp.status_code == 202
        job = wait_for(client, resp.json()["id"], {"done", "failed"})
        assert job["status"] == "failed"
        assert job["error"]
